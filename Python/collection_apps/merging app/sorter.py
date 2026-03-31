import os
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook


# Datei-Auswahl
def select_excel_file(prompt: str) -> str:
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title=prompt,
        filetypes=[("Excel-Dateien", "*.xlsx *.xls")]
    )
    root.destroy()
    if not path:
        raise FileNotFoundError("Keine Datei ausgewählt.")
    return path


# Datei sortieren
def sort_excel(file: str, output_path: str):
    def to_float(value):
        if value is None or value == "":
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        if isinstance(value, str):
            text = value.strip().replace("%", "").replace(".", "").replace(",", ".")
            try:
                num = float(text)
                if "%" in value:
                    return num / 100
                return num
            except ValueError:
                return 0.0

        return 0.0

    wb = load_workbook(file)
    ws = wb.active

    data_start_row = 2
    max_row = ws.max_row

    # Spalte K
    sort_col = 11
    ws.cell(row=1, column=sort_col, value="DB vs Tonnage (relativ)")

    rows_with_score = []

    for row in ws.iter_rows(min_row=data_start_row, max_row=max_row, values_only=False):
        row_values = [ws.cell(row=row[0].row, column=c).value for c in range(1, 11)]

        tonnage_diff = to_float(row[3].value)   # D
        db_diff = to_float(row[9].value)        # J

        # 🔥 Neue robuste Kennzahl
        denominator = abs(tonnage_diff) if abs(tonnage_diff) > 1e-6 else 1e-6
        db_vs_tonnage_rel = (db_diff - tonnage_diff) / denominator

        row_values.append(db_vs_tonnage_rel)

        rows_with_score.append((db_vs_tonnage_rel, row_values))

    # negativ = schlecht → oben
    rows_with_score.sort(key=lambda x: x[0])

    # löschen
    for r in range(data_start_row, max_row + 1):
        for c in range(1, 12):
            ws.cell(row=r, column=c).value = None

    # schreiben
    for new_row_index, (_, row_values) in enumerate(rows_with_score, start=data_start_row):
        for col_index, value in enumerate(row_values, start=1):
            ws.cell(row=new_row_index, column=col_index, value=value)

    wb.save(output_path)



# Main
if __name__ == "__main__":
    try:
        f1 = select_excel_file("Wähle die Excel, die sortiert werden soll")
        default_out = os.path.join(os.getcwd(), "Statistik_Debitoren_2025_sortiert.xlsx")

        root = tk.Tk()
        root.withdraw()
        out = filedialog.asksaveasfilename(
            title="Speichere den sortierten Datensatz",
            defaultextension=".xlsx",
            initialfile="Statistik_Debitoren_2025_sortiert.xlsx",
            filetypes=[("Excel-Dateien", "*.xlsx")]
        )
        root.destroy()

        out = out or default_out

        sort_excel(f1, out)
        print(f"Datei gespeichert unter: {out}")
    except Exception as e:
        print(f"Abbruch: {e}")