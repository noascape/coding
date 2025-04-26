import os
import pandas as pd
import numpy as np
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

desired_columns = ["Nr", "Beschreibung", "Beschreibung2", "Produktgruppencode", "Inventurgruppencode", "Lagerort_24", "Bestand Datum_23", "Bestand Datum_24", "Bestandsveränderung", "Basiseinheit", "Nettogewicht kg/ Einheit_23", "Nettogewicht kg/ Einheit_24", "Nettogewicht kg_23", "Nettogewicht kg_24", "Bewertungspreis € / Einheit_23", "Bewertungspreis € / Einheit_24", "Relative Bewertungspreisveränderung", "Bewertungspreis €_23", "Bewertungspreis €_24", "Absolute Bewertungspreisveränderung"]

#Datei-Auswahl
def select_excel_file(prompt: str) -> str:
    root = tk.Tk(); root.withdraw()
    path = filedialog.askopenfilename(
        title=prompt,
        filetypes=[("Excel-Dateien", "*.xlsx *.xls")]
    )
    root.destroy()
    if not path:
        raise FileNotFoundError("Keine Datei ausgewählt.")
    return path


def merge_excels(file1: str, file2: str, output_path: str):
    # Einlesen mit deutschem Zahlenformat
    cols1 = ["Nr", "Beschreibung_23", "Beschreibung2_23", "Produktgruppencode_23", "Inventurgruppencode_23", "Bestand Datum_23", "Basiseinheit_23", "Nettogewicht kg/ Einheit_23", "Nettogewicht kg_23", "Bewertungspreis € / Einheit_23", "Bewertungspreis €_23"]
    cols2 = ["Nr", "Beschreibung_24", "Beschreibung2_24", "Produktgruppencode_24", "Inventurgruppencode_24", "Lagerort_24", "Bestand Datum_24", "Basiseinheit_24", "Nettogewicht kg/ Einheit_24", "Nettogewicht kg_24", "Bewertungspreis € / Einheit_24", "Bewertungspreis €_24"]
    ds1 = pd.read_excel(file1, skiprows=4, header=None, names=cols1, thousands='.', decimal=',')
    ds2 = pd.read_excel(file2, skiprows=3, header=None, names=cols2, thousands='.', decimal=',')
    df = pd.merge(ds1, ds2, on="Nr", how="outer")

    # Textspalten zusammenführen
    df["Beschreibung"] = df["Beschreibung_24"].combine_first(df["Beschreibung_23"])
    df["Beschreibung2"] = df["Beschreibung2_24"].combine_first(df["Beschreibung2_23"])
    df["Produktgruppencode"] = df["Produktgruppencode_24"].combine_first(df["Produktgruppencode_23"])
    df["Inventurgruppencode"] = df["Inventurgruppencode_24"].combine_first(df["Inventurgruppencode_23"])
    df["Basiseinheit"] = df["Basiseinheit_24"].combine_first(df["Basiseinheit_23"])
    df.drop(columns=["Beschreibung_23","Beschreibung_24", "Beschreibung2_23","Beschreibung2_24", "Produktgruppencode_23","Produktgruppencode_24", "Inventurgruppencode_23","Inventurgruppencode_24", "Basiseinheit_23","Basiseinheit_24"], inplace=True)

    # Custom Spalten (nach Nr gruppiert für die korrekten Berechnungen) 
    sum_23 = df.groupby("Nr")[[
        "Bewertungspreis € / Einheit_23",
        "Bewertungspreis €_23",
        "Bestand Datum_23"
    ]].transform(lambda x: x.sum(min_count=1))
    sum_24 = df.groupby("Nr")[[
        "Bewertungspreis € / Einheit_24",
        "Bewertungspreis €_24",
        "Bestand Datum_24"
    ]].transform(lambda x: x.sum(min_count=1))

    df["Relative Bewertungspreisveränderung"] = (sum_24["Bewertungspreis € / Einheit_24"] - sum_23["Bewertungspreis € / Einheit_23"]) / sum_23["Bewertungspreis € / Einheit_23"]
    df["Absolute Bewertungspreisveränderung"] = (sum_24["Bewertungspreis €_24"] - sum_23["Bewertungspreis €_23"])
    df["Bestandsveränderung"] = (sum_24["Bestand Datum_24"] - sum_23["Bestand Datum_23"])

    # Unendlichkeiten und NaNs ersetzen ohne inplace chaining
    rel = df["Relative Bewertungspreisveränderung"].replace([np.inf, -np.inf], np.nan)
    df["Relative Bewertungspreisveränderung"] = rel

    # Sortieren & Spaltenreihenfolge
    df.sort_values("Absolute Bewertungspreisveränderung", ascending=False, inplace=True)
    df = df[desired_columns]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Zusammengeführt")
        wb = writer.book
        ws = wb["Zusammengeführt"]

        # Legende
        ws.insert_rows(1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
        legend = (
            "Legende: Relative Bewertungspreisveränderung: Rot > 30% | Orange > 20% | Blau = Fehler (Keine Vergleichsdaten vorhanden oder Bewertungspreis pro Einheit = 0)"
        )
        cell = ws.cell(row=1, column=1, value=legend)
        cell.alignment = Alignment(horizontal="center")

        # alle Worksheets sichtbar
        for sheet in wb.worksheets:
            sheet.sheet_state = 'visible'

        # alles oberhalb von A3 freezen
        ws.freeze_panes = 'A3'

        # Breite aller Spalten manuell festlegen (müsste für 24 Zoll passen)
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 6
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 15
        ws.column_dimensions['J'].width = 8
        ws.column_dimensions['K'].width = 8
        ws.column_dimensions['L'].width = 8
        ws.column_dimensions['M'].width = 15
        ws.column_dimensions['N'].width = 15
        ws.column_dimensions['O'].width = 15
        ws.column_dimensions['P'].width = 15
        ws.column_dimensions['Q'].width = 15
        ws.column_dimensions['R'].width = 15
        ws.column_dimensions['S'].width = 15
        ws.column_dimensions['T'].width = 15

        # Richtige Formatierung in 0.00%
        max_col = ws.max_column
        max_row = ws.max_row
        pct_col_idx = None
        for idx_cell, cell in enumerate(ws[2], start=1):
            if cell.value == "Relative Bewertungspreisveränderung":
                pct_col_idx = idx_cell
                break
        if pct_col_idx:
            for row in range(3, max_row+1):
                ws[f"{get_column_letter(pct_col_idx)}{row}"].number_format = '0.00%'

        # Zeilenfärbung
        rel_col = next(
            (c.column_letter for c in ws[2] if c.value == "Relative Bewertungspreisveränderung"),
            None
        )
        if rel_col:
            rng = f"A3:{get_column_letter(max_col)}{max_row}"
            red_fill    = PatternFill(fill_type="solid", start_color="FFFFC7CE", end_color="FFFFC7CE")
            orange_fill = PatternFill(fill_type="solid", start_color="FFFFC000", end_color="FFFFC000")
            blue_fill   = PatternFill(fill_type="solid", start_color="FFCCE5FF", end_color="FFCCE5FF")

            # >30% rot
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f"${rel_col}3>0.30"], fill=red_fill)             # Excel-Formel: gültig für Zeile 3, Excel zieht das automatisch nach unten und passt die Formel entsprechend an
            )
            # >20% orange
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f"${rel_col}3>0.20"], fill=orange_fill)
            )
            # <-30% rot
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f"${rel_col}3<-0.30"], fill=red_fill)
            )
            # <-20% orange
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f"${rel_col}3<-0.20"], fill=orange_fill)
            )
            # NaN (Inf-Ersatz) blau
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f"ISBLANK(${rel_col}3)"], fill=blue_fill)
            )

        # Autofilter für Headerzeile
        ws.auto_filter.ref = f"A2:{get_column_letter(max_col)}{max_row}"

    return df

if __name__ == "__main__":
    try:
        f1 = select_excel_file("Wähle den ersten Excel-Datensatz (z.B. von 2023)")
        f2 = select_excel_file("Wähle den zweiten Excel-Datensatz (z.B. von 2024)")
        default_out = os.path.join(os.getcwd(), "Zusammengeführt.xlsx")
        out = filedialog.asksaveasfilename(
            title="Speichere den zusammengeführten Datensatz",
            defaultextension=".xlsx",
            initialfile="Zusammengeführt.xlsx",
            filetypes=[("Excel-Dateien", "*.xlsx")]
        ) or default_out
        df = merge_excels(f1, f2, out)
        print(f"Datei gespeichert unter: {out}")
    except Exception as e:
        print(f"Abbruch: {e}")
